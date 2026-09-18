import io
import re
from pathlib import Path
from typing import Union
from pypdf import PdfReader
from docx import Document

class TextExtractorError(Exception):
    """Custom exception raised when document extraction fails."""
    pass

def clean_extracted_text(text: str) -> str:
    """Cleans up raw extracted text by removing null bytes, non-printable control characters, and repetitive whitespace."""
    if not text:
        return ""
    # Strip null bytes and non-printable control characters (keep \n, \r, \t)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", " ", text)
    # Replace multiple spaces/tabs with single space
    text = re.sub(r"[ \t]+", " ", text)
    # Replace excessive consecutive newlines with double newline
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def read_pdf_bytes(file_bytes: bytes) -> str:
    """Extracts text from a PDF in-memory byte buffer."""
    text = ""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    except Exception as e:
        raise TextExtractorError(f"Failed to read PDF buffer: {str(e)}")
    
    cleaned = clean_extracted_text(text)
    if not cleaned:
        raise TextExtractorError("PDF appears to be empty or contains scanned images without selectable text.")
    return cleaned

def read_docx_bytes(file_bytes: bytes) -> str:
    """Extracts text from a DOCX in-memory byte buffer including tables."""
    text = ""
    try:
        doc = Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + "\n"
    except Exception as e:
        raise TextExtractorError(f"Failed to read DOCX buffer: {str(e)}")

    cleaned = clean_extracted_text(text)
    if not cleaned:
        raise TextExtractorError("DOCX file contains no readable text.")
    return cleaned

def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    """
    Main extraction dispatcher for FastAPI UploadFile bytes.
    Fixes the Day 4 'raise None' syntax bug and provides clear error handling.
    """
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return read_pdf_bytes(file_bytes)
    elif ext == ".docx":
        return read_docx_bytes(file_bytes)
    else:
        raise TextExtractorError(
            f"Unsupported file format '{ext}'. Only .pdf and .docx files are supported."
        )
