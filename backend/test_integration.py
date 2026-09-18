import asyncio
import io
import docx
from fastapi.testclient import TestClient
from main import app
from database.connection import init_db

client = TestClient(app)

def run_tests():
    print("--- 1. Testing Health Check ---")
    res = client.get("/api/health")
    assert res.status_code == 200, res.text
    print("Health check OK:", res.json())

    print("\n--- 2. Testing HR Registration ---")
    hr_email = "hr_test@example.com"
    hr_pass = "securepassword123"
    res = client.post("/api/auth/register", json={
        "email": hr_email,
        "password": hr_pass,
        "name": "Jane HR",
        "role": "hr"
    })
    if res.status_code == 409: # Already exists from previous run
        res = client.post("/api/auth/login", json={"email": hr_email, "password": hr_pass})
    assert res.status_code in (200, 201), res.text
    hr_token = res.json()["access_token"]
    hr_headers = {"Authorization": f"Bearer {hr_token}"}
    print("HR Registered / Logged In successfully.")

    print("\n--- 3. Testing Duplicate Registration ---")
    res_dup = client.post("/api/auth/register", json={
        "email": hr_email,
        "password": hr_pass,
        "name": "Duplicate Jane",
        "role": "hr"
    })
    assert res_dup.status_code == 409, f"Expected 409, got {res_dup.status_code}"
    print("Duplicate registration correctly rejected with 409 Conflict.")

    print("\n--- 4. Testing /auth/me ---")
    res_me = client.get("/api/auth/me", headers=hr_headers)
    assert res_me.status_code == 200, res_me.text
    user_data = res_me.json()
    assert user_data["email"] == hr_email
    assert user_data["role"] == "hr"
    print("Auth /me verified for HR:", user_data["name"])

    print("\n--- 5. Testing Job Seeker Registration ---")
    seeker_email = "seeker_test@example.com"
    seeker_pass = "seekerpassword123"
    res_seeker = client.post("/api/auth/register", json={
        "email": seeker_email,
        "password": seeker_pass,
        "name": "Alex Seeker",
        "role": "job_seeker"
    })
    if res_seeker.status_code == 409:
        res_seeker = client.post("/api/auth/login", json={"email": seeker_email, "password": seeker_pass})
    assert res_seeker.status_code in (200, 201), res_seeker.text
    seeker_token = res_seeker.json()["access_token"]
    seeker_headers = {"Authorization": f"Bearer {seeker_token}"}
    print("Job Seeker Registered / Logged In successfully.")

    print("\n--- 6. Testing Data Scoping (Empty check) ---")
    res_hr_jobs = client.get("/api/jobs/", headers=hr_headers)
    assert res_hr_jobs.status_code == 200
    res_seeker_jobs = client.get("/api/jobs/", headers=seeker_headers)
    assert res_seeker_jobs.status_code == 200
    print(f"Initial jobs - HR count: {len(res_hr_jobs.json())}, Seeker count: {len(res_seeker_jobs.json())}")

    print("\n--- 7. Testing Sample Resume Generation (.docx) ---")
    doc = docx.Document()
    doc.add_heading("Alex Seeker - Full Stack Engineer", 0)
    doc.add_paragraph("Email: alex.seeker@example.com | Phone: +1-555-0199")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Full stack software developer with 4 years of experience building web applications using React, TypeScript, Python, FastAPI, and PostgreSQL.")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Software Engineer at TechCorp (2022 - Present)")
    doc.add_paragraph("Designed and maintained REST APIs using FastAPI and SQLite/PostgreSQL. Built interactive dashboards using React and Tailwind CSS.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, FastAPI, TypeScript, React, Docker, SQLite, Tailwind CSS, REST APIs")
    
    doc_bytes = io.BytesIO()
    doc.save(doc_bytes)
    doc_bytes.seek(0)
    print("Generated sample .docx resume in memory.")

    print("\n--- 8. Testing Resume Upload ---")
    files = [
        ("files", ("alex_seeker_resume.docx", doc_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    ]
    res_upload = client.post("/api/resumes/upload", headers=hr_headers, files=files)
    assert res_upload.status_code == 200, res_upload.text
    upload_summary = res_upload.json()
    print("Upload summary:", upload_summary)
    assert upload_summary["successful"] >= 1, "Resume should parse successfully"
    parsed_resume_id = upload_summary["resumes"][0]["id"]
    print("Parsed Resume ID:", parsed_resume_id)

    print("\n--- 9. Testing Resume Listing and Scoping ---")
    hr_resumes = client.get("/api/resumes/", headers=hr_headers).json()
    seeker_resumes = client.get("/api/resumes/", headers=seeker_headers).json()
    assert any(r["id"] == parsed_resume_id for r in hr_resumes), "HR should see uploaded resume"
    assert not any(r["id"] == parsed_resume_id for r in seeker_resumes), "Seeker should NOT see HR's uploaded resume (Data Scoping verified!)"
    print(f"Data scoping verified: HR has {len(hr_resumes)} resumes, Seeker has {len(seeker_resumes)} resumes.")

    print("\n--- 10. Testing Job Creation ---")
    raw_jd = """
    Job Title: Senior Full Stack Developer
    Company: Acme Inc.
    Location: Remote
    Required Experience: 3+ years
    Requirements:
    - Strong proficiency in Python, FastAPI, React, and TypeScript.
    - Experience designing relational databases with SQLite or PostgreSQL.
    - Familiarity with modern CSS (Tailwind CSS) and RESTful API design.
    """
    res_job = client.post("/api/jobs/", headers=hr_headers, json={
        "title": "Senior Full Stack Developer",
        "company": "Acme Inc.",
        "raw_text": raw_jd
    })
    assert res_job.status_code == 201, res_job.text
    job_data = res_job.json()
    job_id = job_data["id"]
    print("Created Job ID:", job_id, "Role:", job_data.get("role"))

    print("\n--- 11. Testing Evaluation Pipeline ---")
    res_eval = client.post("/api/match/evaluate", headers=hr_headers, json={
        "job_id": job_id,
        "resume_ids": [parsed_resume_id]
    })
    assert res_eval.status_code == 200, res_eval.text
    eval_data = res_eval.json()
    print("Evaluation results count:", eval_data["total_evaluated"])
    assert len(eval_data["results"]) > 0
    top_candidate = eval_data["results"][0]
    print(f"Candidate: {top_candidate['candidate_name']} | Score: {top_candidate['score']}%")
    print(f"Matching skills: {top_candidate['details']['matching_skills']}")
    print(f"Verdict: {top_candidate['details']['verdict']}")

    print("\n--- 12. Testing CSV Export ---")
    res_csv = client.get(f"/api/match/{job_id}/export", headers=hr_headers)
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers.get("content-type", "")
    assert "Candidate Name" in res_csv.text
    print("CSV Export OK! First 100 chars of CSV:\n", res_csv.text[:100])

    print("\nALL 12 INTEGRATION TESTS PASSED SUCCESSFULLY! [PASSED]")

if __name__ == "__main__":
    asyncio.run(init_db())
    run_tests()
